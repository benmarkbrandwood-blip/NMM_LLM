#!/usr/bin/env python3
"""tools/make_ai_openings_illustrated_docx.py

Generate illustrated Word documents for the Generalist v4 AI's opening preferences,
mirroring the style of docs/opening_plays_illustrated.docx and docs/opening_continuations.docx:

  docs/ai_opening_plays_illustrated.docx       — board image per family at trunk position
  docs/ai_opening_continuations_illustrated.docx — board image in each 3×3 continuation cell

Data source : data/specialist_db_v4.sqlite  (AI self-play, winning lines)
Model       : learned_ai/checkpoints/scaffolded/s_gen_v4/BWDB/best.pt

Run with:
    .venv/bin/python tools/make_ai_openings_illustrated_docx.py
"""
from __future__ import annotations

import collections
import io
import json
import os
import sqlite3
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import matplotlib
matplotlib.use("Agg")
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
import torch
import torch.nn.functional as F
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm

from ai.board_symmetry import canonical_board_str
from game.board import BoardState, MILLS
from game.notation import encode_move, parse_move_string
from game.rules import get_all_legal_moves
from learned_ai.models.scaffolded_encoder import encode_position_with_lookahead
from learned_ai.models.scaffolded_net import ScaffoldedPolicyNet

# ── paths ──────────────────────────────────────────────────────────────────────

DB_PATH    = "data/specialist_db_v4.sqlite"
CKPT_PATH  = "learned_ai/checkpoints/scaffolded/s_gen_v4/BWDB/best.pt"
OUT_PLAYS  = "docs/ai_opening_plays_illustrated.docx"
OUT_CONTS  = "docs/ai_opening_continuations_illustrated.docx"

# ── family definitions ─────────────────────────────────────────────────────────

FAMILY_NAMES = {
    ("d6", "f4"): "Corner Cross",
    ("d6", "d7"): "Diagonal Setup",
    ("f4", "b4"): "Counter Flank",
    ("f4", "d2"): "Center Column",
    ("d2", "b4"): "Column Hold",
    ("d6", "b4"): "Wing Anchor",
    ("d6", "d2"): "Double Column",
    ("f4", "b6"): "Column Strike",
    ("f4", "d5"): "Center Rush",
    ("b4", "d5"): "Reverse Center",
    ("d6", "a7"): "Corner Setup",
    ("b4", "d2"): "Column Mirror",
}

# ── board geometry ─────────────────────────────────────────────────────────────

_POS_XY: dict[str, tuple[float, float]] = {
    "a1": (0, 0), "d1": (3, 0), "g1": (6, 0),
    "a4": (0, 3), "b4": (1, 3), "c4": (2, 3),
    "e4": (4, 3), "f4": (5, 3), "g4": (6, 3),
    "a7": (0, 6), "d7": (3, 6), "g7": (6, 6),
    "b2": (1, 1), "d2": (3, 1), "f2": (5, 1),
    "b6": (1, 5), "d6": (3, 5), "f6": (5, 5),
    "c3": (2, 2), "d3": (3, 2), "e3": (4, 2),
    "c5": (2, 4), "d5": (3, 4), "e5": (4, 4),
}
_EDGES: list[tuple[str, str]] = [
    ("a7", "d7"), ("d7", "g7"), ("g7", "g4"), ("g4", "g1"),
    ("g1", "d1"), ("d1", "a1"), ("a1", "a4"), ("a4", "a7"),
    ("b6", "d6"), ("d6", "f6"), ("f6", "f4"), ("f4", "f2"),
    ("f2", "d2"), ("d2", "b2"), ("b2", "b4"), ("b4", "b6"),
    ("c5", "d5"), ("d5", "e5"), ("e5", "e4"), ("e4", "e3"),
    ("e3", "d3"), ("d3", "c3"), ("c3", "c4"), ("c4", "c5"),
    ("a4", "b4"), ("b4", "c4"),
    ("d7", "d6"), ("d6", "d5"),
    ("g4", "f4"), ("f4", "e4"),
    ("d1", "d2"), ("d2", "d3"),
    ("a7", "a4"), ("a4", "a1"),
    ("g7", "g4"), ("g4", "g1"),
]

_BOARD_BG   = "#D4920A"
_LINE_COLOR = "black"
_W_FACE     = "#f2ede0"
_W_EDGE     = "#555555"
_B_FACE     = "#1e1a2e"
_B_EDGE     = "#9090b0"
_MILL_COLOR = "#CC0000"
_PIECE_R    = 0.42
_MILL_R     = 0.18

# ── board rendering ────────────────────────────────────────────────────────────

def _detect_mills(positions: dict[str, str]) -> dict[str, list[tuple]]:
    """Return {color: [list of mill triples]} from a positions dict."""
    result: dict[str, list] = {"W": [], "B": []}
    for triple in MILLS:
        colors = [positions.get(p, "") for p in triple]
        if colors[0] and colors[0] == colors[1] == colors[2]:
            result[colors[0]].append(tuple(triple))
    return result


def _board_png(
    positions: dict[str, str],
    label: str,
    last_move_to: str | None = None,
    dpi: int = 110,
) -> bytes:
    """Render the board with both W (cream) and B (dark) pieces. Returns PNG bytes."""
    fig, ax = plt.subplots(figsize=(2.2, 2.5))
    fig.patch.set_facecolor(_BOARD_BG)
    ax.set_facecolor(_BOARD_BG)

    for p1, p2 in _EDGES:
        x1, y1 = _POS_XY[p1]; x2, y2 = _POS_XY[p2]
        ax.plot([x1, x2], [y1, y2], color=_LINE_COLOR, linewidth=1.8, zorder=1)
    for pos, (x, y) in _POS_XY.items():
        ax.plot(x, y, ".", color=_LINE_COLOR, markersize=3, zorder=2)

    mills = _detect_mills(positions)
    mill_squares: set[str] = set()
    for triples in mills.values():
        for triple in triples:
            mill_squares.update(triple)
    for pos in mill_squares:
        x, y = _POS_XY[pos]
        ax.add_patch(mpatches.Circle((x, y), _MILL_R, color=_MILL_COLOR, zorder=3))

    # Highlight last move destination with a ring
    if last_move_to and last_move_to in _POS_XY:
        lx, ly = _POS_XY[last_move_to]
        ax.add_patch(mpatches.Circle(
            (lx, ly), _PIECE_R + 0.12,
            facecolor="none", edgecolor="#FFD700", linewidth=2.0, zorder=4,
        ))

    for pos, color in positions.items():
        if color not in ("W", "B"):
            continue
        x, y = _POS_XY[pos]
        face = _W_FACE if color == "W" else _B_FACE
        edge = _W_EDGE if color == "W" else _B_EDGE
        ax.add_patch(mpatches.Circle(
            (x, y), _PIECE_R,
            facecolor=face, edgecolor=edge, linewidth=1.2, zorder=5,
        ))

    ax.set_xlim(-0.85, 6.85)
    ax.set_ylim(-1.1, 7.1)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.text(3, -0.8, label, ha="center", va="center", fontsize=8,
            fontweight="bold", color="black", transform=ax.transData)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor=_BOARD_BG)
    plt.close(fig)
    buf.seek(0)
    return buf.read()

# ── model helpers ──────────────────────────────────────────────────────────────

def load_model(ckpt_path: str, device: str = "cpu"):
    ckpt = torch.load(ckpt_path, map_location=device, weights_only=False)
    cfg  = ckpt.get("model_config", {})
    model = ScaffoldedPolicyNet.from_config(cfg)
    model.load_state_dict(ckpt["model"])
    model.eval()
    difficulty = ckpt.get("difficulty", "?")
    game_count = ckpt.get("game_count", 0)
    return model, device, difficulty, game_count


def replay_moves(move_strings: list[str]) -> BoardState:
    board = BoardState.new_game()
    for ms in move_strings:
        board = board.apply_move(parse_move_string(ms))
    return board


def score_legal_moves(model, board: BoardState, device: str, top_n: int = 3) -> list[dict]:
    enc = encode_position_with_lookahead(
        board, board.turn,
        sentinel_advisor=None, db=None, value_net=None, lookahead_advisor=None,
    )
    if enc is None or not enc.legal_moves:
        return []
    feat_t = torch.tensor(enc.feat_matrix, dtype=torch.float32)
    with torch.no_grad():
        logits = model.policy_logits(feat_t)
        probs  = F.softmax(logits, dim=-1)
    order = probs.argsort(descending=True)
    return [
        {**enc.legal_moves[i.item()], "_prob": float(probs[i].item())}
        for i in order[:top_n]
    ]


def mv_notation(mv: dict, color: str, player_num: int, board: BoardState) -> str:
    """Format a move as W5:d6 or B5:f4 (per-player move number)."""
    raw = encode_move(mv, board.phase)
    return f"{color}{player_num}:{raw}"


def format_trunk(moves: list[str]) -> str:
    board = BoardState.new_game()
    parts = []
    w_num = b_num = 0
    for ms in moves:
        mv    = parse_move_string(ms)
        color = board.turn
        raw   = encode_move(mv, board.phase)
        if color == "W":
            w_num += 1
            parts.append(f"W{w_num}:{raw}")
        else:
            b_num += 1
            parts.append(f"B{b_num}:{raw}")
        board = board.apply_move(mv)
    lines = []
    for j in range(0, len(parts), 2):
        lines.append("  ".join(parts[j:j+2]))
    return "\n".join(lines)

# ── DB extraction ──────────────────────────────────────────────────────────────

def _canon_key_2ply(moves: list[str]) -> str:
    """D4-canonical key for the board position after the first 2 moves."""
    board = replay_moves(moves[:2])
    fen   = board.to_fen_string()
    parts = fen.split("|")
    canon, _ = canonical_board_str(parts[0])
    return f"{canon}|{parts[1]}|{parts[2]}|{parts[3]}"


def extract_families(db_path: str, min_diff: int = 0):
    conn = sqlite3.connect(db_path)
    c    = conn.cursor()
    # Use diff-filtered table if available and requested
    has_diff_table = c.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name='winning_lines_by_diff'"
    ).fetchone() is not None
    if min_diff > 0 and has_diff_table:
        c.execute(
            'SELECT move_seq FROM winning_lines_by_diff '
            'WHERE result IN ("W", "draw") AND diff_level >= ?',
            (min_diff,)
        )
    else:
        c.execute('SELECT move_seq FROM winning_lines WHERE result IN ("W", "draw")')
    rows = c.fetchall()
    conn.close()

    families:     dict[str, list]                       = collections.defaultdict(list)
    raw_counters: dict[str, collections.Counter]        = collections.defaultdict(collections.Counter)

    for (seq,) in rows:
        moves = json.loads(seq)
        if len(moves) >= 2:
            canon_key = _canon_key_2ply(moves)
            families[canon_key].append(moves)
            raw_counters[canon_key][tuple(moves[:2])] += 1

    # Require 11 families with enough clean-placement trunks; scan top candidates
    MIN_PLACE_GAMES = 2   # minimum placement-only games needed for a valid trunk
    ranked_all = sorted(families.items(), key=lambda x: -len(x[1]))
    result = []
    rank = 0
    for canon_key, games in ranked_all:
        if len(result) >= 11:
            break
        raw_key = raw_counters[canon_key].most_common(1)[0][0]
        name    = FAMILY_NAMES.get(raw_key, f"{raw_key[0]}-{raw_key[1]}")

        # Only use games whose first 8 moves are ALL placements (no slide notation).
        # winning_lines stores single-player sequences; slide moves in the prefix
        # corrupt the board when replayed on an alternating-turn board.
        place_games = [
            g for g in games
            if len(g) >= 8 and all("-" not in m for m in g[:8])
        ]
        if len(place_games) < MIN_PLACE_GAMES:
            continue  # skip family — not enough clean-placement trunk data

        rank += 1
        code = f"A{rank:02d}"
        prefix8: collections.Counter = collections.Counter()
        for g in place_games:
            prefix8[tuple(g[:8])] += 1
        trunk8, trunk_cnt = prefix8.most_common(1)[0]
        assert len(trunk8) == 8, f"{code}: trunk has {len(trunk8)} moves (expected 8)"
        result.append((code, name, raw_key, list(trunk8), len(games)))
    return result, len(rows)

# ── docx helpers ───────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_colour: str) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_colour)
    tcPr.append(shd)


def _set_no_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{side}")
                border.set(qn("w:val"), "none")
                tcBorders.append(border)
            tcPr.append(tcBorders)


def _para(doc, text: str, bold: bool = False, size: int = 10):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p

# ── plays illustrated builder ──────────────────────────────────────────────────

def build_plays_illustrated(families, total_games: int, difficulty, game_count: int) -> Document:
    doc = Document()
    doc.core_properties.title = "NMM AI Opening Plays — Illustrated (gen v4)"

    doc.add_heading("NMM AI Opening Plays — gen v4 (Illustrated)", level=0)
    _para(doc,
        f"Top 11 opening families · Generalist v4 · difficulty {difficulty}/20 · "
        f"{game_count:,} games · {total_games:,} winning/draw lines from {DB_PATH}.",
        size=10)
    _para(doc,
        "Board shows trunk position after 8 plies. ● cream = White, ● dark = Black. "
        "Gold ring = last move. Red dot = mill.",
        size=9)
    doc.add_page_break()

    for code, name, key, trunk, game_count_fam in families:
        doc.add_heading(f"{code}  {name}", level=1)
        _para(doc, f"8-ply  ·  DB: {game_count_fam} games  ·  seed: {key[0]}-{key[1]}", size=9)

        trunk_board = replay_moves(trunk)
        last_to     = parse_move_string(trunk[-1]).get("to") if trunk else None
        png = _board_png(trunk_board.positions, f"{code} {name}", last_move_to=last_to, dpi=120)

        # Two-column table: image | move sequence + notes
        tbl = doc.add_table(rows=1, cols=2)
        _set_no_borders(tbl)
        tbl.columns[0].width = Inches(3.0)
        tbl.columns[1].width = Inches(3.5)

        img_cell  = tbl.cell(0, 0)
        text_cell = tbl.cell(0, 1)

        img_para = img_cell.paragraphs[0]
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_para.add_run().add_picture(io.BytesIO(png), width=Inches(2.8))

        tp = text_cell.paragraphs[0]
        tp.add_run("Moves:\n").bold = True
        tp.add_run(format_trunk(trunk)).font.size = Pt(9)
        text_cell.add_paragraph()

        doc.add_paragraph()

    return doc

# ── continuations illustrated builder ─────────────────────────────────────────

def build_continuation_cells_with_boards(
    model, trunk_moves: list[str], device: str, trunk_ply: int, code: str
) -> list[tuple[str, bytes]]:
    """Return 9 (notation, png_bytes) tuples for the 3×3 table.

    Continuation is 4 moves deep (ply 9–12 = W5/B5/W6/B6 per-player).
    """
    trunk_board = replay_moves(trunk_moves)
    side_a = trunk_board.turn                  # colour that moves at ply trunk+1
    side_b = "B" if side_a == "W" else "W"

    # Per-player move numbers beyond the trunk
    # W goes at odd global plies (1,3,5,...), B at even plies (2,4,6,...).
    w_in_trunk = (trunk_ply + 1) // 2   # = 4 for trunk_ply=8
    b_in_trunk = trunk_ply // 2         # = 4 for trunk_ply=8

    if side_a == "W":
        num_a = w_in_trunk + 1   # W5
        num_b = b_in_trunk + 1   # B5
        num_c = w_in_trunk + 2   # W6
        num_d = b_in_trunk + 2   # B6
    else:
        num_a = b_in_trunk + 1   # B5
        num_b = w_in_trunk + 1   # W5
        num_c = b_in_trunk + 2   # B6
        num_d = w_in_trunk + 2   # W6

    top_a = score_legal_moves(model, trunk_board, device, top_n=3)

    cells: list[tuple[str, bytes]] = []
    for mv_a in top_a:
        board_a   = trunk_board.apply_move(mv_a)
        nota_a    = mv_notation(mv_a, side_a, num_a, trunk_board)
        last_to_a = mv_a.get("to")

        top_b = score_legal_moves(model, board_a, device, top_n=3)
        if not top_b:
            png = _board_png(board_a.positions, f"{code} {nota_a}", last_move_to=last_to_a, dpi=90)
            for _ in range(3):
                cells.append((f"{nota_a}\n(terminal)\n★ gen v4", png))
            continue

        for mv_b in top_b:
            board_b   = board_a.apply_move(mv_b)
            nota_b    = mv_notation(mv_b, side_b, num_b, board_a)
            last_to_b = mv_b.get("to")

            top_c = score_legal_moves(model, board_b, device, top_n=1)
            if not top_c:
                notation = f"{nota_a}  {nota_b}\n★ gen v4"
                label    = f"{nota_a}\n{nota_b}"
                png = _board_png(board_b.positions, label, last_move_to=last_to_b, dpi=90)
                cells.append((notation, png))
                continue

            board_c   = board_b.apply_move(top_c[0])
            nota_c    = mv_notation(top_c[0], side_a, num_c, board_b)
            last_to_c = top_c[0].get("to")

            top_d = score_legal_moves(model, board_c, device, top_n=1)
            if top_d:
                board_d   = board_c.apply_move(top_d[0])
                nota_d    = mv_notation(top_d[0], side_b, num_d, board_c)
                last_to_d = top_d[0].get("to")
                notation  = f"{nota_a}  {nota_b}\n{nota_c}  {nota_d}\n★ gen v4"
                label     = f"{nota_a}\n{nota_d}"
                png = _board_png(board_d.positions, label, last_move_to=last_to_d, dpi=90)
            else:
                notation = f"{nota_a}  {nota_b}\n{nota_c}\n★ gen v4"
                label    = f"{nota_a}\n{nota_c}"
                png = _board_png(board_c.positions, label, last_move_to=last_to_c, dpi=90)
            cells.append((notation, png))

    while len(cells) < 9:
        empty_png = _board_png({}, "", dpi=90)
        cells.append(("", empty_png))
    return cells[:9]


def build_continuations_illustrated(families, model, device: str, difficulty) -> Document:
    doc = Document()
    doc.core_properties.title = "NMM AI Opening Continuations — Illustrated (gen v4)"

    doc.add_heading("NMM AI Opening Continuations — gen v4 (Illustrated)", level=0)
    _para(doc,
        f"Top 9 continuations per family from W5/B5 to W6/B6 (3 W choices × 3 B responses, 4 moves deep). "
        f"★ gen v4 = Generalist v4 policy (difficulty {difficulty}/20). "
        f"Gold ring = last move. Red dot = mill.", size=10)
    doc.add_page_break()

    for code, name, key, trunk, game_count_fam in families:
        doc.add_heading(f"{code}  {name}", level=1)
        _para(doc, f"Trunk (8-ply):  {format_trunk(trunk)}", size=9)
        _para(doc, f"★ = Generalist v4 · DB: {game_count_fam} games", size=9)
        doc.add_paragraph()

        print(f"  Generating continuations for {code} {name} ...", flush=True)
        cells = build_continuation_cells_with_boards(model, trunk, device, trunk_ply=8, code=code)

        tbl = doc.add_table(rows=3, cols=3)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        BG = "EEF4FF"
        for ri in range(3):
            for ci in range(3):
                idx  = ri * 3 + ci
                cell = tbl.cell(ri, ci)
                _set_cell_bg(cell, BG)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                notation, png = cells[idx]

                # Board image
                img_p = cell.paragraphs[0]
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(io.BytesIO(png), width=Inches(1.5))

                # Move text below image
                txt_p = cell.add_paragraph()
                txt_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = txt_p.add_run(notation)
                run.font.size = Pt(8)

        doc.add_paragraph()

    return doc

# ── main ───────────────────────────────────────────────────────────────────────

def main() -> None:
    import argparse
    ap = argparse.ArgumentParser(description="Generate NMM AI opening docs")
    ap.add_argument("--db",       default=DB_PATH,  help="SpecialistDB path")
    ap.add_argument("--ckpt",     default=CKPT_PATH, help="Model checkpoint path")
    ap.add_argument("--min-diff", type=int, default=0,
                    help="Minimum difficulty level for winning_lines_by_diff filter "
                         "(0 = use aggregate winning_lines, gen v4 style)")
    args = ap.parse_args()

    os.makedirs("docs", exist_ok=True)

    print(f"Loading DB (min_diff={args.min_diff}) …")
    families, total_games = extract_families(args.db, min_diff=args.min_diff)
    print(f"  {len(families)} families, {total_games} total games")

    print("Loading gen v4 model …")
    model, device, difficulty, game_count = load_model(CKPT_PATH)
    print(f"  difficulty={difficulty}, game_count={game_count:,}")

    print("Building ai_opening_plays_illustrated.docx …")
    plays_doc = build_plays_illustrated(families, total_games, difficulty, game_count)
    plays_doc.save(OUT_PLAYS)
    print(f"  Saved: {OUT_PLAYS}")

    print("Building ai_opening_continuations_illustrated.docx …")
    conts_doc = build_continuations_illustrated(families, model, device, difficulty)
    conts_doc.save(OUT_CONTS)
    print(f"  Saved: {OUT_CONTS}")

    print("Done.")


if __name__ == "__main__":
    main()
